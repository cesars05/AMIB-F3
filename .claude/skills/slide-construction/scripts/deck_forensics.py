"""Measure the real formatting of a .pptx: every shape's geometry, position, fill,
line, fonts and every native chart's XML parameters.

Usage:
    python deck_forensics.py deck.pptx [--out forensics.json] [--summary]

Output JSON: {deck, slide_w, slide_h, slides:[{n, layout, hidden, shapes:[...], charts:[...]}]}
Positions are given in inches (x,y,w,h) AND as fractions of the slide (fx,fy) so
decks of different sizes (10x7.5, 13.33x7.5) can be compared.
--summary prints the aggregate profile (shape-geometry counts, font sizes by
character volume, fill colors, chart parameters) to compare against
references/style-systems.md.
"""
import argparse
import collections
import json

from lxml import etree
from pptx import Presentation

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
IN = 914400


def X(el, path):
    """Namespaced XPath that also works on python-pptx oxml elements."""
    return etree.XPath(path, namespaces=NS)(el)


def inch(v):
    return None if v is None else round(v / IN, 3)


def color_of(el):
    """Hex or scheme name of the first solidFill child of el (spPr / ln / tcPr)."""
    if el is None:
        return None
    f = el.find("a:solidFill", NS)
    if f is None:
        return "none" if el.find("a:noFill", NS) is not None else None
    if not len(f):
        return None
    c = f[0]
    tag, val = etree.QName(c).localname, c.get("val")
    mods = {etree.QName(m).localname: m.get("val") for m in c}
    base = "#" + val if tag == "srgbClr" else f"{tag}:{val}"
    return base + ("" if not mods else " " + json.dumps(mods))


def line_of(spPr):
    ln = spPr.find("a:ln", NS) if spPr is not None else None
    if ln is None:
        return None
    dash = ln.find("a:prstDash", NS)
    return {"w_pt": round(int(ln.get("w")) / 12700, 2) if ln.get("w") else None,
            "color": color_of(ln), "dash": dash.get("val") if dash is not None else None}


def is_group(sh):
    return sh.shape_type is not None and "GROUP" in str(sh.shape_type)


def shape_record(sh, sw, shh):
    el = sh._element
    spPr = el.find(".//p:spPr", NS)
    geom = None
    if spPr is not None:
        pg = spPr.find("a:prstGeom", NS)
        geom = pg.get("prst") if pg is not None else ("custGeom" if spPr.find("a:custGeom", NS) is not None else None)
    rec = {"name": sh.name,
           "type": str(sh.shape_type).split(".")[-1].split(" ")[0] if sh.shape_type else "?",
           "geom": geom,
           "x": inch(sh.left), "y": inch(sh.top), "w": inch(sh.width), "h": inch(sh.height),
           "fx": round(sh.left / sw, 3) if sh.left is not None else None,
           "fy": round(sh.top / shh, 3) if sh.top is not None else None,
           "fill": color_of(spPr), "line": line_of(spPr)}
    if spPr is not None:
        adj = [g.get("fmla") for g in spPr.findall(".//a:avLst/a:gd", NS)]
        if adj:
            rec["adj"] = adj
    if sh.has_text_frame and sh.text_frame.text.strip():
        runs = []
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                f = r.font
                col = None
                try:
                    if f.color and f.color.type is not None:
                        col = str(f.color.rgb) if f.color.type == 1 else f"theme:{f.color.theme_color}"
                except Exception:
                    pass
                runs.append({"t": r.text[:60], "sz": f.size.pt if f.size else None, "b": f.bold,
                             "i": f.italic, "font": f.name, "color": col,
                             "align": str(p.alignment).split(".")[-1].split(" ")[0] if p.alignment else None,
                             "lvl": p.level})
        rec["text"] = sh.text_frame.text[:200]
        rec["runs"] = runs[:25]
        bodyPr = el.find(".//a:bodyPr", NS)
        if bodyPr is not None:
            rec["anchor"] = bodyPr.get("anchor")
    if getattr(sh, "has_table", False) and sh.has_table:
        t = sh.table
        fills = []
        for r in range(min(len(t.rows), 8)):
            fills.append([color_of(t.cell(r, c)._tc.find("a:tcPr", NS)) for c in range(len(t.columns))])
        rec["table"] = {"rows": len(t.rows), "cols": len(t.columns),
                        "row_h": [inch(r.height) for r in t.rows][:40],
                        "col_w": [inch(c.width) for c in t.columns],
                        "cells": [[t.cell(r, c).text[:30] for c in range(len(t.columns))] for r in range(min(len(t.rows), 6))],
                        "fills": fills}
    if is_group(sh):
        rec["children"] = [shape_record(c, sw, shh) for c in sh.shapes]
    return rec


def chart_record(sh):
    x = sh.chart._chartSpace
    q = lambda p: X(x, p)
    plots = []
    for pa in q("//c:plotArea/*"):
        tag = etree.QName(pa).localname
        if not tag.endswith("Chart"):
            continue
        d = {"type": tag}
        for k in ("barDir", "grouping", "gapWidth", "overlap", "holeSize", "varyColors"):
            e = pa.find(f"c:{k}", NS)
            if e is not None:
                d[k] = e.get("val")
        sers = []
        for s in pa.findall("c:ser", NS):
            tx = X(s, ".//c:tx//c:v/text()")
            spPr = s.find("c:spPr", NS)
            dl = s.find("c:dLbls", NS)
            lab = None
            if dl is not None:
                g = lambda k: (dl.find(f"c:{k}", NS).get("val") if dl.find(f"c:{k}", NS) is not None else None)
                sz = X(dl, ".//a:defRPr/@sz")
                nf = dl.find("c:numFmt", NS)
                lab = {"pos": g("dLblPos"), "showVal": g("showVal"), "sz": int(sz[0]) / 100 if sz else None,
                       "bold": (X(dl, ".//a:defRPr/@b") or [None])[0],
                       "color": color_of(dl.find(".//a:defRPr", NS)),
                       "numFmt": nf.get("formatCode") if nf is not None else None}
            sers.append({"name": tx[0] if tx else None, "fill": color_of(spPr), "line": line_of(spPr),
                         "pts": len(X(s, "./c:val//c:pt")) or len(X(s, "./c:yVal//c:pt")),
                         "dpt_overrides": len(s.findall("c:dPt", NS)), "labels": lab,
                         "marker": (X(s, "./c:marker/c:symbol/@val") or [None])[0]})
        d["series"] = sers
        plots.append(d)
    axes = []
    for ax in q("//c:plotArea/c:catAx|//c:plotArea/c:valAx|//c:plotArea/c:dateAx"):
        dele = ax.find("c:delete", NS)
        nf = ax.find("c:numFmt", NS)
        sz = X(ax, ".//a:defRPr/@sz")
        axes.append({"kind": etree.QName(ax).localname,
                     "deleted": dele.get("val") if dele is not None else None,
                     "gridlines": ax.find("c:majorGridlines", NS) is not None,
                     "numFmt": nf.get("formatCode") if nf is not None else None,
                     "tick_sz": int(sz[0]) / 100 if sz else None,
                     "line": line_of(ax.find("c:spPr", NS))})
    leg = q("//c:legend/c:legendPos/@val")
    return {"name": sh.name, "x": inch(sh.left), "y": inch(sh.top), "w": inch(sh.width), "h": inch(sh.height),
            "plots": plots, "axes": axes, "legend": leg[0] if leg else None,
            "title": bool(q("//c:chart/c:title")) and not q("//c:chart/c:autoTitleDeleted[@val='1']")}


def walk(shapes):
    for s in shapes:
        yield s
        if is_group(s):
            yield from walk(s.shapes)


def analyse(path):
    prs = Presentation(path)
    sw, shh = prs.slide_width, prs.slide_height
    out = {"deck": path, "slide_w": inch(sw), "slide_h": inch(shh), "slides": []}
    for i, sl in enumerate(prs.slides, 1):
        rec = {"n": i, "layout": sl.slide_layout.name, "hidden": sl._element.get("show") == "0",
               "shapes": [shape_record(s, sw, shh) for s in sl.shapes], "charts": []}
        for s in walk(sl.shapes):
            if getattr(s, "has_chart", False) and s.has_chart:
                try:
                    rec["charts"].append(chart_record(s))
                except Exception as e:  # keep going on exotic charts
                    rec["charts"].append({"name": s.name, "error": str(e)})
        out["slides"].append(rec)
    return out


def summary(d):
    geoms, sizes, fills, fonts, dashes = (collections.Counter() for _ in range(5))

    def visit(shs):
        for s in shs:
            if s.get("geom"):
                geoms[s["geom"]] += 1
            if s.get("fill") and s["fill"] != "none":
                fills[s["fill"].split(" ")[0]] += 1
            if s.get("line") and s["line"].get("dash"):
                dashes[s["line"]["dash"]] += 1
            for r in s.get("runs", []):
                if r["sz"]:
                    sizes[r["sz"]] += len(r["t"])
                if r["font"]:
                    fonts[r["font"]] += len(r["t"])
            visit(s.get("children", []))

    for sl in d["slides"]:
        visit(sl["shapes"])
    print(f"# {d['deck']}  ({d['slide_w']}x{d['slide_h']} in, {len(d['slides'])} slides)")
    print("geometries:", geoms.most_common(25))
    print("font sizes by chars:", sorted(sizes.items(), key=lambda kv: -kv[1])[:14])
    print("fonts:", fonts.most_common(5))
    print("fills:", fills.most_common(18))
    print("dashes:", dashes.most_common(5))
    ch = [c for sl in d["slides"] for c in sl["charts"]]
    print("charts:", len(ch), collections.Counter(p["type"] + "/" + str(p.get("grouping"))
                                                  for c in ch for p in c.get("plots", [])).most_common())
    plots = [p for c in ch for p in c.get("plots", [])]
    print("  gap/overlap:", collections.Counter((p.get("gapWidth"), p.get("overlap")) for p in plots if "gapWidth" in p).most_common(6))
    print("  label sz/pos:", collections.Counter((s["labels"] or {}).get("sz") for p in plots for s in p["series"]).most_common(5),
          collections.Counter((s["labels"] or {}).get("pos") for p in plots for s in p["series"]).most_common(5))
    print("  series fills:", collections.Counter((s["fill"] or (s["line"] or {}).get("color") or "auto").split(" ")[0]
                                                for p in plots for s in p["series"]).most_common(10))
    print("  line widths:", collections.Counter((s["line"] or {}).get("w_pt") for p in plots for s in p["series"]
                                               if p["type"] in ("lineChart", "scatterChart")).most_common(4))
    axes = [a for c in ch for a in c.get("axes", [])]
    print("  axes (kind,deleted,grid):", collections.Counter((a["kind"], a["deleted"], a["gridlines"]) for a in axes).most_common(6))
    print("  legend:", collections.Counter(c.get("legend") for c in ch).most_common())


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("--out")
    ap.add_argument("--summary", action="store_true")
    a = ap.parse_args()
    data = analyse(a.deck)
    if a.out:
        with open(a.out, "w", encoding="utf-8") as fh:
            json.dump(data, fh, indent=1, ensure_ascii=False)
    if a.summary or not a.out:
        summary(data)
