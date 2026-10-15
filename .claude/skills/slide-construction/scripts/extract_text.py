"""Dump all slide text (shapes, groups, tables, chart series, notes) as markdown with
'## Slide N' markers - the input format for extract_numbers.py.

    python extract_text.py deck.pptx > deck_content.md
"""
import sys

from pptx import Presentation


def shape_text(sh, out):
    if getattr(sh, "shape_type", None) is not None and "GROUP" in str(sh.shape_type):
        for c in sh.shapes:
            shape_text(c, out)
        return
    if getattr(sh, "has_text_frame", False) and sh.has_text_frame and sh.text_frame.text.strip():
        out.append(sh.text_frame.text.strip())
    if getattr(sh, "has_table", False) and sh.has_table:
        for r in sh.table.rows:
            out.append(" | ".join(c.text.strip() for c in r.cells))
    if getattr(sh, "has_chart", False) and sh.has_chart:
        ch = sh.chart
        try:
            cats = list(ch.plots[0].categories)
        except Exception:
            cats = []
        for pl in ch.plots:
            for se in pl.series:
                if se.name and se.name.startswith("_"):
                    continue  # helper series (totals, offsets)
                try:
                    vals = list(se.values)
                except Exception:
                    vals = []
                pairs = ", ".join(f"{c}: {v}" for c, v in zip(cats, vals))
                out.append(f"[chart series {se.name}] {pairs}")


def main(path):
    prs = Presentation(path)
    for i, sl in enumerate(prs.slides, 1):
        out = []
        for sh in sl.shapes:
            shape_text(sh, out)
        if sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip():
            out.append("[notes] " + sl.notes_slide.notes_text_frame.text.strip())
        print(f"## Slide {i}")
        print("\n".join(out))
        print()


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main(sys.argv[1])
